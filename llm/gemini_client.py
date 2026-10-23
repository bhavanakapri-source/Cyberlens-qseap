"""
Gemini API client for CyberLens.

Uses:
    google-genai
    Gemini model configured in .env

Also records Gemini token usage in:

    RAG/logs/gemini_token_usage.xlsx
"""

import os
import time
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
from google import genai
from openpyxl import Workbook, load_workbook


# ============================================================
# ENVIRONMENT
# ============================================================

load_dotenv()

GEMINI_API_KEY = os.getenv(
    "GEMINI_API_KEY"
)

GEMINI_MODEL = os.getenv(
    "GEMINI_MODEL",
    "gemini-3.5-flash-lite",
)


# ============================================================
# PROJECT PATH
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

LOG_DIR = PROJECT_ROOT / "logs"

TOKEN_LOG_FILE = (
    LOG_DIR / "gemini_token_usage.xlsx"
)


# ============================================================
# CLIENT
# ============================================================

_client = None


def get_client():
    """
    Create Gemini client lazily.

    This prevents the API client from being created
    every time a function is called.
    """

    global _client

    if _client is None:

        if not GEMINI_API_KEY:
            raise ValueError(
                "GEMINI_API_KEY is not set. "
                "Add it to the .env file."
            )

        _client = genai.Client(
            api_key=GEMINI_API_KEY
        )

    return _client


# ============================================================
# EXTRACT USER QUERY
# ============================================================

def _extract_user_query(prompt: str) -> str:
    """
    Extract the original user question from the prompt.

    The prompt_builder places the user question between:

        USER QUESTION
        ...
        CYBERLENS RETRIEVED KNOWLEDGE

    This allows us to log the actual user query without
    changing the existing FastAPI route.
    """

    start_marker = (
        "============================================================\n"
        "USER QUESTION\n"
        "============================================================\n"
    )

    end_marker = (
        "\n============================================================\n"
        "CYBERLENS RETRIEVED KNOWLEDGE\n"
        "============================================================"
    )

    try:

        start_index = prompt.index(
            start_marker
        )

        start_index += len(
            start_marker
        )

        end_index = prompt.index(
            end_marker,
            start_index
        )

        return prompt[
            start_index:end_index
        ].strip()

    except ValueError:

        return ""


# ============================================================
# CREATE TOKEN LOG FILE
# ============================================================

def _create_token_log_file():
    """
    Create the Excel token usage file if it does not exist.
    """

    LOG_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    if TOKEN_LOG_FILE.exists():
        return

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Gemini Usage"

    worksheet.append([
        "Timestamp",
        "User Query",
        "Model",
        "Prompt Tokens",
        "Output Tokens",
        "Thought Tokens",
        "Cached Tokens",
        "Total Tokens",
        "Latency Seconds",
        "Status",
        "Error",
    ])

    workbook.save(
        TOKEN_LOG_FILE
    )


# ============================================================
# LOG TOKEN USAGE
# ============================================================

def _log_token_usage(
    user_query: str,
    prompt_tokens: int,
    output_tokens: int,
    thought_tokens: int,
    cached_tokens: int,
    total_tokens: int,
    latency_seconds: float,
    status: str,
    error: str = "",
):
    """
    Append one Gemini request to the Excel token log.
    """

    _create_token_log_file()

    workbook = load_workbook(
        TOKEN_LOG_FILE
    )

    worksheet = workbook[
        "Gemini Usage"
    ]

    worksheet.append([
        datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        ),
        user_query,
        GEMINI_MODEL,
        prompt_tokens,
        output_tokens,
        thought_tokens,
        cached_tokens,
        total_tokens,
        round(
            latency_seconds,
            3
        ),
        status,
        error,
    ])

    workbook.save(
        TOKEN_LOG_FILE
    )


# ============================================================
# GENERATE ANSWER
# ============================================================

def generate_answer(
    prompt: str,
) -> str:
    """
    Send the grounded prompt to Gemini and return
    the generated answer.

    Token usage is automatically logged to Excel.
    """

    client = get_client()

    user_query = _extract_user_query(
        prompt
    )

    start_time = time.perf_counter()

    try:

        response = client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
        )

        latency_seconds = (
            time.perf_counter()
            - start_time
        )

        # ----------------------------------------------------
        # GEMINI USAGE METADATA
        # ----------------------------------------------------

        usage = getattr(
            response,
            "usage_metadata",
            None,
        )

        prompt_tokens = 0
        output_tokens = 0
        thought_tokens = 0
        cached_tokens = 0
        total_tokens = 0

        if usage:

            prompt_tokens = (
                getattr(
                    usage,
                    "prompt_token_count",
                    0,
                )
                or 0
            )

            output_tokens = (
                getattr(
                    usage,
                    "candidates_token_count",
                    0,
                )
                or 0
            )

            thought_tokens = (
                getattr(
                    usage,
                    "thoughts_token_count",
                    0,
                )
                or 0
            )

            cached_tokens = (
                getattr(
                    usage,
                    "cached_content_token_count",
                    0,
                )
                or 0
            )

            total_tokens = (
                getattr(
                    usage,
                    "total_token_count",
                    0,
                )
                or 0
            )

        # ----------------------------------------------------
        # EMPTY RESPONSE
        # ----------------------------------------------------

        if not response.text:

            _log_token_usage(
                user_query=user_query,
                prompt_tokens=prompt_tokens,
                output_tokens=output_tokens,
                thought_tokens=thought_tokens,
                cached_tokens=cached_tokens,
                total_tokens=total_tokens,
                latency_seconds=latency_seconds,
                status="error",
                error=(
                    "Gemini returned an empty response."
                ),
            )

            raise RuntimeError(
                "Gemini returned an empty response."
            )

        answer = response.text.strip()

        # ----------------------------------------------------
        # SAVE SUCCESSFUL REQUEST
        # ----------------------------------------------------

        _log_token_usage(
            user_query=user_query,
            prompt_tokens=prompt_tokens,
            output_tokens=output_tokens,
            thought_tokens=thought_tokens,
            cached_tokens=cached_tokens,
            total_tokens=total_tokens,
            latency_seconds=latency_seconds,
            status="success",
        )

        # ----------------------------------------------------
        # TERMINAL LOG
        # ----------------------------------------------------

        print(
            "\n"
            "============================================================\n"
            "GEMINI TOKEN USAGE\n"
            "============================================================\n"
            f"Model          : {GEMINI_MODEL}\n"
            f"User Query     : {user_query}\n"
            f"Prompt Tokens  : {prompt_tokens}\n"
            f"Output Tokens  : {output_tokens}\n"
            f"Thought Tokens : {thought_tokens}\n"
            f"Cached Tokens  : {cached_tokens}\n"
            f"Total Tokens   : {total_tokens}\n"
            f"Latency        : {latency_seconds:.3f} seconds\n"
            f"Excel Log      : {TOKEN_LOG_FILE}\n"
            "============================================================\n"
        )

        return answer

    except Exception as exc:

        latency_seconds = (
            time.perf_counter()
            - start_time
        )

        _log_token_usage(
            user_query=user_query,
            prompt_tokens=0,
            output_tokens=0,
            thought_tokens=0,
            cached_tokens=0,
            total_tokens=0,
            latency_seconds=latency_seconds,
            status="error",
            error=str(exc),
        )

        raise


# ============================================================
# SIMPLE HEALTH CHECK
# ============================================================

def test_connection() -> str:
    """
    Simple Gemini connection test.

    This is only a health check and is not written
    to the chatbot token usage Excel file.
    """

    client = get_client()

    response = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=(
            "Reply with exactly: "
            "CyberLens Gemini connection successful."
        ),
    )

    if not response.text:
        raise RuntimeError(
            "Gemini returned an empty response."
        )

    return response.text.strip()