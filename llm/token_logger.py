from pathlib import Path
from datetime import datetime

from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent

LOG_DIR = PROJECT_ROOT / "logs"
LOG_FILE = LOG_DIR / "gemini_token_usage.xlsx"


HEADERS = [
    "Timestamp",
    "Conversation ID",
    "User Query",
    "Model",
    "Input Tokens",
    "Output Tokens",
    "Thoughts Tokens",
    "Total Tokens",
    "Latency (seconds)",
    "Status",
    "Error",
]


def log_token_usage(
    user_query,
    conversation_id,
    model,
    usage_metadata,
    latency_seconds,
    status="success",
    error="",
):
    LOG_DIR.mkdir(parents=True, exist_ok=True)

    if LOG_FILE.exists():
        workbook = load_workbook(LOG_FILE)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Token Usage"
        sheet.append(HEADERS)

    prompt_tokens = getattr(
        usage_metadata,
        "prompt_token_count",
        0,
    ) or 0

    output_tokens = getattr(
        usage_metadata,
        "candidates_token_count",
        0,
    ) or 0

    thoughts_tokens = getattr(
        usage_metadata,
        "thoughts_token_count",
        0,
    ) or 0

    total_tokens = getattr(
        usage_metadata,
        "total_token_count",
        0,
    ) or 0

    sheet.append(
        [
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
            conversation_id,
            user_query,
            model,
            prompt_tokens,
            output_tokens,
            thoughts_tokens,
            total_tokens,
            round(latency_seconds, 3),
            status,
            error,
        ]
    )

    workbook.save(LOG_FILE)

    print(
        f"[TOKEN LOG] "
        f"Input={prompt_tokens}, "
        f"Output={output_tokens}, "
        f"Total={total_tokens}"
    )